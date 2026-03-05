from pptx import Presentation
import os


class PresentationBuilder:

    def __init__(self, client_name, data):
        self.client_name = client_name
        self.data = data
        self.prs = Presentation()

    # --------------------------------------------------

    def add_title_slide(self):

        slide_layout = self.prs.slide_layouts[0]
        slide = self.prs.slides.add_slide(slide_layout)

        slide.shapes.title.text = f"{self.client_name} Growth Strategy"

        subtitle = slide.placeholders[1]
        subtitle.text = "AI Marketing Consultant System"

    # --------------------------------------------------

    def add_summary_slide(self):

        slide_layout = self.prs.slide_layouts[1]
        slide = self.prs.slides.add_slide(slide_layout)

        slide.shapes.title.text = "Executive Summary"

        tf = slide.placeholders[1].text = self.data["recommendation"]

    # --------------------------------------------------

    def add_business_diagnosis(self):

        slide_layout = self.prs.slide_layouts[1]
        slide = self.prs.slides.add_slide(slide_layout)

        slide.shapes.title.text = "Business Diagnosis"

        content = f"""
Rating: {self.data['rating']}
Reviews: {self.data['reviews']}
Social Status: {self.data['social_status']}
Core Problem: {self.data['core_problem']}
"""

        slide.placeholders[1].text = content

    # --------------------------------------------------

    def add_strategy_slide(self):

        slide_layout = self.prs.slide_layouts[1]
        slide = self.prs.slides.add_slide(slide_layout)

        slide.shapes.title.text = "Growth Strategy"

        content = f"""
Primary Strategy: {self.data['primary_lever']}
Campaign Framework: {self.data['campaign_type']}
Content Strategy: {self.data['offer_strategy']}
Growth Focus: {self.data['sales_angle']}
"""

        slide.placeholders[1].text = content

    # --------------------------------------------------

    def add_action_plan(self):

        slide_layout = self.prs.slide_layouts[1]
        slide = self.prs.slides.add_slide(slide_layout)

        slide.shapes.title.text = "90 Day Action Plan"

        content = f"""
Month 1: {self.data['month1']}
Month 2: {self.data['month2']}
Month 3: {self.data['month3']}
"""

        slide.placeholders[1].text = content

    # --------------------------------------------------

    def build(self):

        self.add_title_slide()
        self.add_summary_slide()
        self.add_business_diagnosis()
        self.add_strategy_slide()
        self.add_action_plan()

    # --------------------------------------------------

    def save(self):

        self.build()

        os.makedirs("clients", exist_ok=True)

        path = f"clients/{self.client_name}_growth_strategy.pptx"

        self.prs.save(path)

        return path